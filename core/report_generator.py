"""报告生成 — ReportGenerator，使用 Jinja2 模板生成多种格式报告"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Dict, List, Optional
from datetime import datetime

import openpyxl
from jinja2 import Environment, FileSystemLoader, BaseLoader

from domain.report import DepartmentReport, ReportSection
from domain.analysis import FullAnalysisReport, AnalysisResult
from infra.exceptions import ReportError
from infra.event_bus import REPORT_GENERATED, EventBus
from infra.logger import get_logger

logger = get_logger()


class ReportGenerator:
    """报告生成器，使用 Jinja2 模板生成多种格式报告。

    报告结构:
        - 执行摘要
        - 各部门详细分析
        - 综合建议
        - 附录
    """

    def __init__(
        self,
        templates_dir: str = "templates",
        event_bus: Optional[EventBus] = None,
    ) -> None:
        """初始化报告生成器。

        Args:
            templates_dir: Jinja2 模板目录
            event_bus: 事件总线
        """
        self._event_bus = event_bus or EventBus()
        self._env = Environment(
            loader=FileSystemLoader(templates_dir),
            autoescape=False,
            trim_blocks=True,
            lstrip_blocks=True,
        )

        # 注册自定义过滤器
        self._env.filters["format_number"] = self._format_number
        self._env.filters["format_percent"] = self._format_percent
        self._env.filters["format_datetime"] = self._format_datetime
        self._env.filters["truncate"] = self._truncate

    def generate_markdown(self, report: FullAnalysisReport) -> str:
        """生成 Markdown 格式报告。

        Args:
            report: 分析报告

        Returns:
            Markdown 报告内容
        """
        template = self._env.get_template("report_markdown.j2")
        return template.render(report=report, generated_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

    def generate_excel(self, report: FullAnalysisReport, filepath: str) -> str:
        """生成 Excel 格式报告。

        Args:
            report: 分析报告
            filepath: 输出文件路径

        Returns:
            实际输出的文件路径
        """
        path = Path(filepath)
        path.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()

        # Sheet 1: 执行摘要
        ws_summary = wb.active
        ws_summary.title = "执行摘要"
        ws_summary.append(["AI业务调研提效分析报告", "", "", ""])
        ws_summary.append([f"生成时间: {report.generated_at or datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", "", "", ""])
        ws_summary.append([f"分析部门数: {len(report.departments)}", "", "", ""])
        ws_summary.append([f"总机会点数: {report.total_opportunities}", "", "", ""])
        ws_summary.append([f"高优先级机会点: {report.high_priority_opportunities}", "", "", ""])
        ws_summary.append(["", "", "", ""])
        ws_summary.append(["总体总结", "", "", ""])
        ws_summary.append([report.overall_summary, "", "", ""])

        # Sheet 2: 各部门详细分析
        ws_dept = wb.create_sheet("各部门分析")
        ws_dept.append(["部门", "痛点数量", "机会点数量", "置信度", "总结"])
        for dept in report.departments:
            ws_dept.append([
                dept.department,
                len(dept.pain_points),
                len(dept.opportunities),
                f"{dept.confidence:.0%}",
                dept.summary,
            ])

        # Sheet 3: 机会点明细
        ws_opp = wb.create_sheet("机会点明细")
        ws_opp.append(["部门", "业务环节", "当前流程", "AI建议", "预期提效", "难度", "优先级", "ROI", "周期"])
        for dept in report.departments:
            for opp in dept.opportunities:
                ws_opp.append([
                    opp.department,
                    opp.process,
                    opp.current_process,
                    opp.ai_suggestion,
                    opp.expected_efficiency_gain,
                    opp.difficulty,
                    opp.priority,
                    opp.expected_roi or "",
                    opp.timeline or "",
                ])

        # Sheet 4: 综合建议
        ws_sug = wb.create_sheet("综合建议")
        ws_sug.append(["AI业务调研提效 - 综合建议", "", "", ""])
        ws_sug.append([self._generate_suggestions(report), "", "", ""])

        wb.save(str(path))
        wb.close()

        logger.info("Excel 报告已生成: %s", path)
        return str(path)

    def generate_json(self, report: FullAnalysisReport, filepath: str) -> str:
        """生成 JSON 格式报告。

        Args:
            report: 分析报告
            filepath: 输出文件路径

        Returns:
            实际输出的文件路径
        """
        path = Path(filepath)
        path.parent.mkdir(parents=True, exist_ok=True)

        # 将报告转为可序列化的字典
        data = {
            "report_metadata": {
                "generated_at": report.generated_at or datetime.now().isoformat(),
                "total_departments": len(report.departments),
                "total_opportunities": report.total_opportunities,
                "high_priority_opportunities": report.high_priority_opportunities,
            },
            "overall_summary": report.overall_summary,
            "departments": [],
        }

        for dept in report.departments:
            dept_data = {
                "department": dept.department,
                "pain_points": dept.pain_points,
                "opportunities": [],
                "summary": dept.summary,
                "confidence": dept.confidence,
            }
            for opp in dept.opportunities:
                dept_data["opportunities"].append({
                    "process": opp.process,
                    "current_process": opp.current_process,
                    "ai_suggestion": opp.ai_suggestion,
                    "expected_efficiency_gain": opp.expected_efficiency_gain,
                    "difficulty": opp.difficulty,
                    "priority": opp.priority,
                    "expected_roi": opp.expected_roi,
                    "timeline": opp.timeline,
                })
            data["departments"].append(dept_data)

        path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        logger.info("JSON 报告已生成: %s", path)
        return str(path)

    def generate_summary(self, report: FullAnalysisReport) -> str:
        """生成纯文本摘要。

        Args:
            report: 分析报告

        Returns:
            摘要文本
        """
        lines: list[str] = [
            "=" * 60,
            "AI业务调研提效分析报告",
            "=" * 60,
            f"\n生成时间: {report.generated_at or datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            f"\n分析部门数: {len(report.departments)}",
            f"总机会点数: {report.total_opportunities}",
            f"高优先级机会点: {report.high_priority_opportunities}",
            f"\n{'=' * 60}\n",
            "总体总结",
            "-" * 40,
            report.overall_summary,
            f"\n{'=' * 60}\n",
        ]

        for dept in report.departments:
            lines.append(f"\n{dept.department}")
            lines.append("-" * 40)

            if dept.pain_points:
                lines.append("痛点:")
                for i, pt in enumerate(dept.pain_points, 1):
                    lines.append(f"  {i}. {pt}")

            if dept.opportunities:
                lines.append("\nAI提效机会:")
                for i, opp in enumerate(dept.opportunities, 1):
                    lines.append(f"  {i}. [{opp.difficulty}] {opp.process} - {opp.ai_suggestion}")
                    lines.append(f"     提效: {opp.expected_efficiency_gain}, 优先级: {opp.priority}/5")

            lines.append(f"\n总结: {dept.summary}")
            lines.append(f"置信度: {dept.confidence:.0%}")

        lines.append(f"\n{'=' * 60}")
        return "\n".join(lines)

    def generate_department_report(self, analysis_result: AnalysisResult) -> DepartmentReport:
        """将 AnalysisResult 转换为 DepartmentReport。

        Args:
            analysis_result: 单个部门的分析结果

        Returns:
            DepartmentReport 对象
        """
        sections: list[ReportSection] = []

        # 执行摘要
        sections.append(ReportSection(
            title="执行摘要",
            content=analysis_result.summary,
        ))

        # 痛点分析
        sections.append(ReportSection(
            title="痛点分析",
            content="\n".join(f"- {pt}" for pt in analysis_result.pain_points),
        ))

        # 机会点
        opp_content = []
        for opp in analysis_result.opportunities:
            opp_content.append(
                f"**{opp.process}** - {opp.ai_suggestion} (提效: {opp.expected_efficiency_gain}, 难度: {opp.difficulty})"
            )
        sections.append(ReportSection(
            title="AI提效机会",
            content="\n".join(opp_content),
        ))

        # 建议
        recommendation = self._generate_department_recommendation(analysis_result)
        sections.append(ReportSection(
            title="建议方案",
            content=recommendation,
        ))

        return DepartmentReport(
            department=analysis_result.department,
            executive_summary=analysis_result.summary,
            pain_points=analysis_result.pain_points,
            opportunities=[opp.ai_suggestion for opp in analysis_result.opportunities],
            recommendation=recommendation,
            sections=sections,
        )

    def _generate_suggestions(self, report: FullAnalysisReport) -> str:
        """生成综合建议。"""
        lines: list[str] = [
            "1. **短期优先**: 聚焦高优先级、低难度的提效机会，快速见效。",
            "2. **中期规划**: 制定跨部门协同方案，整合共享AI能力。",
            "3. **长期战略**: 建立AI治理框架，持续优化AI应用效果。",
            "4. **组织保障**: 设立AI推进小组，负责技术选型、实施和效果追踪。",
            "5. **人才建设**: 培养AI应用能力，加强业务培训。",
        ]
        return "\n".join(lines)

    def _generate_department_recommendation(self, result: AnalysisResult) -> str:
        """生成部门建议。"""
        if not result.opportunities:
            return "当前部门暂无明显的AI提效机会，建议持续跟踪业务变化。"

        high_priority = [opp for opp in result.opportunities if opp.priority >= 4]
        medium_priority = [opp for opp in result.opportunities if 2 <= opp.priority < 4]

        parts: list[str] = []
        if high_priority:
            parts.append("短期优先推进:")
            for opp in high_priority:
                parts.append(f"- {opp.ai_suggestion} (难度: {opp.difficulty}, 提效: {opp.expected_efficiency_gain})")

        if medium_priority:
            parts.append("中期规划:")
            for opp in medium_priority:
                parts.append(f"- {opp.ai_suggestion} (难度: {opp.difficulty})")

        return "\n".join(parts)

    def _format_number(self, value: Any) -> str:
        """格式化数字。"""
        try:
            num = float(value)
            return f"{num:,.0f}"
        except (ValueError, TypeError):
            return str(value)

    def _format_percent(self, value: Any) -> str:
        """格式化百分比。"""
        try:
            num = float(value)
            return f"{num:.0%}"
        except (ValueError, TypeError):
            return str(value)

    def _format_datetime(self, dt: Any) -> str:
        """格式化日期时间。"""
        if isinstance(dt, datetime):
            return dt.strftime("%Y-%m-%d %H:%M:%S")
        return str(dt)

    def _truncate(self, text: str, length: int = 100) -> str:
        """截断文本。"""
        if len(text) <= length:
            return text
        return text[:length] + "..."
