"""数据采集 — DataCollector，支持交互式采集、结构化采集和多格式文件解析"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

from docx import Document
from openpyxl import load_workbook

from domain.template import BusinessTemplate, QuestionType
from infra.exceptions import CollectionError
from infra.event_bus import DATA_COLLECTED, EventBus
from infra.logger import get_logger

logger = get_logger()


def _parse_markdown_data(content: str) -> Dict[str, Any]:
    """从 Markdown 内容解析调研数据。

    约定格式:
        ### Q1 问题文本
        答案: 文本内容

        ### Q2 选择题
        选项: A, B, C
        答案: A
    """
    answers: dict[str, str] = {}
    current_qid: Optional[str] = None

    for line in content.strip().splitlines():
        stripped = line.strip()

        # 问题编号
        m_qid = re.match(r"^###?\s+(Q\d+)\s+(.+)$", stripped)
        if m_qid:
            current_qid = m_qid.group(1)
            continue

        # 答案
        m_answer = re.match(r"^答案:\s*(.+)$", stripped)
        if m_answer and current_qid:
            answers[current_qid] = m_answer.group(1).strip()
            continue

    return {"answers": answers, "raw": content}


def _parse_txt_data(content: str) -> Dict[str, Any]:
    """从 TXT 内容解析调研数据（使用 Markdown 格式解析）。"""
    return _parse_markdown_data(content)


def _parse_json_data(content: str) -> Dict[str, Any]:
    """从 JSON 内容解析调研数据。"""
    try:
        data = json.loads(content)
        if "answers" in data:
            return data
        # 如果 JSON 直接包含键值对
        return {"answers": data, "raw": content}
    except json.JSONDecodeError as e:
        raise CollectionError(f"JSON 解析失败: {e}")


def _parse_docx_data(file_path: Path) -> Dict[str, Any]:
    """从 DOCX 文档解析调研数据。"""
    doc = Document(str(file_path))
    answers: dict[str, str] = {}
    current_qid: Optional[str] = None

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        m_qid = re.match(r"^###?\s+(Q\d+)\s+(.+)$", text)
        if m_qid:
            current_qid = m_qid.group(1)
            continue

        m_answer = re.match(r"^答案:\s*(.+)$", text)
        if m_answer and current_qid:
            answers[current_qid] = m_answer.group(1).strip()

    return {"answers": answers, "raw": "\n".join(p.text for p in doc.paragraphs)}


def _parse_excel_data(file_path: Path) -> Dict[str, Any]:
    """从 Excel 文件解析调研数据。

    约定: 第一列为问题编号，第二列为答案。
    """
    wb = load_workbook(str(file_path), read_only=True)
    ws = wb.active
    answers: dict[str, str] = {}

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if len(row) >= 2 and row[0] and row[1]:
            qid = str(row[0]).strip()
            answer = str(row[1]).strip()
            if qid.startswith("Q"):
                answers[qid] = answer

    wb.close()
    return {"answers": answers}


class DataCollector:
    """数据采集器，支持交互式采集、结构化采集和文件格式解析。

    支持的文件格式:
        - Markdown (.md)
        - TXT (.txt)
        - DOCX (.docx)
        - Excel (.xlsx)
        - JSON (.json)
    """

    def __init__(self, event_bus: Optional[EventBus] = None) -> None:
        self._event_bus = event_bus or EventBus()

    def collect_interactive(self, template: BusinessTemplate) -> Dict[str, str]:
        """交互式采集调研数据。

        Args:
            template: 调研模板

        Returns:
            答案字典 {问题编号: 答案文本}
        """
        answers: dict[str, str] = {}

        for question in template.questions:
            prompt = f"\n[{question.question_type.value}] {question.text}"
            if question.description:
                prompt += f"\n  ({question.description})"

            if question.options:
                options_str = ", ".join(f"[{i + 1}] {o}" for i, o in enumerate(question.options))
                prompt += f"\n  选项: {options_str}"

            prompt += "\n  请输入答案: "
            answer = input(prompt)

            if question.required and not answer.strip():
                raise CollectionError(f"问题 {question.id} 为必填项，不能为空")

            answers[question.id] = answer.strip()

        self._event_bus.publish(DATA_COLLECTED, {"template": template.title, "answers": answers})
        return answers

    def collect_structured(self, data: Dict[str, Any]) -> Dict[str, str]:
        """结构化采集调研数据（从已有数据提取）。

        Args:
            data: 结构化数据，支持 {"answers": {...}} 或直接 {"Q1": "答案1"}

        Returns:
            标准化答案字典
        """
        if "answers" in data:
            answers = data["answers"]
        else:
            # 过滤出以 Q 开头的键
            answers = {k: v for k, v in data.items() if k.startswith("Q")}

        if not answers:
            raise CollectionError("未找到有效的调研数据")

        self._event_bus.publish(DATA_COLLECTED, {"data": answers})
        return answers

    def validate(self, data: Dict[str, str], template: BusinessTemplate) -> Dict[str, Any]:
        """验证采集的数据是否符合模板要求。

        Args:
            data: 答案数据
            template: 调研模板

        Returns:
            验证结果 {"valid": bool, "errors": [...]}
        """
        errors: list[str] = []
        warnings: list[str] = []

        # 检查必填项
        for question in template.questions:
            if question.required:
                if question.id not in data or not data[question.id].strip():
                    errors.append(f"必填问题 {question.id} 为空")

        # 检查选项型问题
        for question in template.questions:
            if question.options:
                answer = data.get(question.id, "").strip()
                if answer and answer not in question.options:
                    # 尝试匹配选项编号
                    if answer.isdigit():
                        idx = int(answer) - 1
                        if 0 <= idx < len(question.options):
                            continue  # 编号匹配成功
                    warnings.append(f"问题 {question.id} 的答案 '{answer}' 不在选项列表中")

        # 检查问题编号
        known_ids = {q.id for q in template.questions}
        for qid in data:
            if qid not in known_ids:
                warnings.append(f"未知问题编号: {qid}")

        is_valid = len(errors) == 0
        return {"valid": is_valid, "errors": errors, "warnings": warnings}

    def export(self, data: Dict[str, Any], filepath: str) -> str:
        """导出调研数据到文件。

        Args:
            data: 数据字典
            filepath: 输出文件路径

        Returns:
            实际输出的文件路径
        """
        path = Path(filepath)
        path.parent.mkdir(parents=True, exist_ok=True)

        ext = path.suffix.lower()
        if ext in (".json", ""):
            path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        elif ext == ".md":
            path.write_text(self._format_markdown(data), encoding="utf-8")
        elif ext == ".txt":
            path.write_text(self._format_markdown(data), encoding="utf-8")
        elif ext == ".docx":
            self._write_docx(data, path)
        else:
            path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

        logger.info("数据已导出到: %s", path)
        return str(path)

    def parse_file(self, file_path: str) -> Dict[str, Any]:
        """解析多种格式的调研数据文件。

        Args:
            file_path: 文件路径

        Returns:
            解析后的数据字典
        """
        path = Path(file_path)
        if not path.exists():
            raise CollectionError(f"文件不存在: {file_path}")

        ext = path.suffix.lower()
        try:
            if ext == ".md":
                content = path.read_text(encoding="utf-8")
                return _parse_markdown_data(content)
            elif ext == ".txt":
                content = path.read_text(encoding="utf-8")
                return _parse_txt_data(content)
            elif ext == ".json":
                content = path.read_text(encoding="utf-8")
                return _parse_json_data(content)
            elif ext == ".docx":
                return _parse_docx_data(path)
            elif ext in (".xlsx", ".xls"):
                return _parse_excel_data(path)
            else:
                raise CollectionError(f"不支持的文件格式: {ext}")
        except CollectionError:
            raise
        except Exception as e:
            raise CollectionError(f"文件解析失败 {file_path}: {e}")

    def parse_directory(self, directory: str) -> Dict[str, Dict[str, Any]]:
        """解析目录下所有调研数据文件。

        Args:
            directory: 目录路径

        Returns:
            {文件名: 解析数据} 字典
        """
        dir_path = Path(directory)
        if not dir_path.is_dir():
            raise CollectionError(f"目录不存在: {directory}")

        results: dict[str, dict[str, Any]] = {}
        for file_path in sorted(dir_path.glob("**/*")):
            if file_path.is_file() and file_path.suffix.lower() in (".md", ".txt", ".json", ".docx", ".xlsx"):
                try:
                    results[file_path.name] = self.parse_file(str(file_path))
                except CollectionError as e:
                    logger.warning("跳过文件 %s: %s", file_path.name, e)
        return results

    def _format_markdown(self, data: Dict[str, Any]) -> str:
        """将数据格式化为 Markdown 文本。"""
        lines: list[str] = ["# 调研数据\n"]

        if "answers" in data:
            answers = data["answers"]
        else:
            answers = data

        for qid, answer in answers.items():
            lines.append(f"### {qid}\n")
            lines.append(f"答案: {answer}\n")

        return "\n".join(lines)

    def _write_docx(self, data: Dict[str, Any], path: Path) -> None:
        """将数据写入 DOCX 文档。"""
        doc = Document()

        if "answers" in data:
            answers = data["answers"]
        else:
            answers = data

        doc.add_heading("调研数据", level=1)
        for qid, answer in answers.items():
            doc.add_paragraph(f"### {qid}", style="Heading 2")
            doc.add_paragraph(f"答案: {answer}")

        doc.save(str(path))
