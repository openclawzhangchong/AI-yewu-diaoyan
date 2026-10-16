"""core 单元测试 — TemplateManager / DataCollector / AIAnalysisEngine。"""

from __future__ import annotations

import json

import pytest

from core.ai_engine import AIAnalysisEngine
from core.data_collector import DataCollector
from core.template_manager import TemplateManager
from domain.analysis import AnalysisResult
from domain.template import BusinessTemplate, Question, QuestionType
from infra.exceptions import AnalysisError, CollectionError, LLMError, TemplateError


@pytest.fixture
def data_collector(event_bus):
    return DataCollector(event_bus=event_bus)


def _dept_data(department, answers=None, template=None):
    return {"department": department, "answers": answers or {}, "template": template}


# ---------------------------------------------------------------------------
# TemplateManager
# ---------------------------------------------------------------------------


class TestTemplateManager:
    def test_load_all_returns_five_department_templates(self, template_manager):
        titles = set(template_manager.list_templates())
        expected = {
            "运营部业务调研模板",
            "客服部业务调研模板",
            "设计部业务调研模板",
            "仓储部业务调研模板",
            "财务部业务调研模板",
        }
        assert titles == expected
        for tpl in template_manager.load_all():
            assert len(tpl.questions) == 5, f"{tpl.title} 应有 5 个问题"

    def test_get_template(self, template_manager):
        tpl = template_manager.get_template("运营部业务调研模板")
        assert tpl is not None
        assert tpl.department_hint == "运营部"
        assert [q.id for q in tpl.questions] == ["Q1", "Q2", "Q3", "Q4", "Q5"]
        assert template_manager.get_template("不存在的模板") is None

    def test_register_requires_title(self):
        tm = TemplateManager()
        tpl = BusinessTemplate(
            title="",
            description="x",
            questions=[Question("Q1", "问题?", QuestionType.PAIN_POINT)],
        )
        with pytest.raises(TemplateError):
            tm.register(tpl)

    def test_register_requires_questions(self):
        tm = TemplateManager()
        with pytest.raises(TemplateError):
            tm.register(BusinessTemplate(title="T", description="D", questions=[]))

    def test_register_and_get(self):
        tm = TemplateManager()
        tpl = BusinessTemplate(
            title="新模板",
            description="desc",
            questions=[Question("Q1", "问题?", QuestionType.PAIN_POINT)],
        )
        tm.register(tpl)
        assert tm.get_template("新模板") is tpl
        assert tm.list_templates() == ["新模板"]
        assert tm.remove_template("新模板") is True
        assert tm.get_template("新模板") is None
        assert tm.remove_template("新模板") is False

    def test_get_by_department(self, template_manager):
        matches = template_manager.get_by_department("财务部")
        assert len(matches) == 1
        assert matches[0].department_hint == "财务部"
        assert template_manager.get_by_department("不存在的部门") == []

    def test_load_missing_dir_returns_empty(self, tmp_path):
        tm = TemplateManager(template_dir=str(tmp_path / "nope"))
        assert tm.load_all() == []

    def test_json_parse_contract(self, tmp_path):
        """json 模板解析：type 必须是 QuestionType 合法值，否则该文件被跳过。"""
        good = tmp_path / "good.json"
        good.write_text(
            json.dumps(
                {
                    "title": "好模板",
                    "description": "d",
                    "department_hint": "测试部",
                    "questions": [
                        {"id": "Q1", "text": "t", "type": "pain_point", "required": True}
                    ],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        bad_type = tmp_path / "bad_type.json"
        bad_type.write_text(
            json.dumps(
                {
                    "title": "坏类型",
                    "questions": [
                        {"id": "Q1", "text": "t", "type": "not_a_type"}
                    ],
                }
            ),
            encoding="utf-8",
        )
        bad_json = tmp_path / "broken.json"
        bad_json.write_text("{oops", encoding="utf-8")

        tm = TemplateManager(template_dir=str(tmp_path))
        loaded = tm.load_all()
        titles = [t.title for t in loaded]
        assert titles == ["好模板"]  # 两个坏文件被跳过且不抛错
        assert tm.get_template("好模板").questions[0].question_type is QuestionType.PAIN_POINT

    def test_load_all_reloads_from_disk(self, repo_templates_dir):
        """load_all 每次都从目录重新扫描，不会重复注册。"""
        tm = TemplateManager(str(repo_templates_dir))
        first = tm.load_all()
        second = tm.load_all()
        assert [t.title for t in first] == [t.title for t in second]


# ---------------------------------------------------------------------------
# DataCollector
# ---------------------------------------------------------------------------


class TestDataCollectorParseFile:
    def test_parse_markdown(self, data_collector, md_fixture_path):
        parsed = data_collector.parse_file(str(md_fixture_path))
        assert set(parsed) == {"answers", "raw"}
        answers = parsed["answers"]
        assert answers["Q1"].startswith("客户进线")
        assert answers["Q4"] == "每天大量（>100单）"
        assert len(answers) == 5

    def test_parse_txt(self, data_collector, txt_fixture_path):
        parsed = data_collector.parse_file(str(txt_fixture_path))
        assert parsed["answers"]["Q3"]
        assert len(parsed["answers"]) == 5

    def test_parse_json(self, data_collector, json_fixture_path):
        parsed = data_collector.parse_file(str(json_fixture_path))
        assert parsed["answers"]["Q1"]
        assert parsed["answers"]["Q5"] == "每天 100-1000件"
        assert "raw" not in parsed or isinstance(parsed.get("raw"), str)

    def test_parse_missing_file(self, data_collector, tmp_path):
        with pytest.raises(CollectionError):
            data_collector.parse_file(str(tmp_path / "missing.md"))

    def test_parse_unsupported_ext(self, data_collector, tmp_path):
        f = tmp_path / "data.pdf"
        f.write_bytes(b"%PDF")
        with pytest.raises(CollectionError, match="不支持的文件格式"):
            data_collector.parse_file(str(f))

    def test_parse_invalid_json_raises_collection_error(self, data_collector, tmp_path):
        f = tmp_path / "bad.json"
        f.write_text("{broken", encoding="utf-8")
        with pytest.raises(CollectionError):
            data_collector.parse_file(str(f))

    def test_parse_docx(self, data_collector, tmp_path):
        from docx import Document

        f = tmp_path / "doc.docx"
        doc = Document()
        doc.add_paragraph("### Q1 问题")
        doc.add_paragraph("答案: 流程描述")
        doc.add_paragraph("### Q2 频次")
        doc.add_paragraph("答案: 每天")
        doc.save(str(f))
        parsed = data_collector.parse_file(str(f))
        assert parsed["answers"] == {"Q1": "流程描述", "Q2": "每天"}

    def test_parse_excel(self, data_collector, tmp_path):
        import openpyxl

        f = tmp_path / "doc.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["问题编号", "答案"])
        ws.append(["Q1", "流程描述"])
        ws.append(["Q2", "每天"])
        ws.append(["非Q", "忽略"])
        wb.save(str(f))
        parsed = data_collector.parse_file(str(f))
        assert parsed["answers"] == {"Q1": "流程描述", "Q2": "每天"}

    def test_parse_directory(self, data_collector, tmp_path):
        (tmp_path / "a.md").write_text("### Q1 问题\n答案: X\n", encoding="utf-8")
        (tmp_path / "b.json").write_text('{"answers": {"Q1": "Y"}}', encoding="utf-8")
        (tmp_path / "c.bad").write_text("junk", encoding="utf-8")
        results = data_collector.parse_directory(str(tmp_path))
        assert set(results) == {"a.md", "b.json"}
        assert results["a.md"]["answers"] == {"Q1": "X"}
        assert results["b.json"]["answers"] == {"Q1": "Y"}


class TestDataCollectorValidate:
    def test_validate_pass(self, data_collector, yunying_template, yunying_answers):
        result = data_collector.validate(yunying_answers, yunying_template)
        assert result["valid"] is True
        assert result["errors"] == []
        assert result["warnings"] == []

    def test_validate_kefu_fixture_passes(self, data_collector, md_fixture_path, template_manager):
        """tests/fixtures 样例 md 对客服部模板应通过 validate。"""
        parsed = data_collector.parse_file(str(md_fixture_path))
        tpl = template_manager.get_template("客服部业务调研模板")
        result = data_collector.validate(parsed["answers"], tpl)
        assert result["valid"] is True
        assert result["errors"] == []

    def test_validate_missing_required(self, data_collector, yunying_template, yunying_answers):
        data = {k: v for k, v in yunying_answers.items() if k != "Q4"}  # 删掉必填 Q4
        result = data_collector.validate(data, yunying_template)
        assert result["valid"] is False
        assert any("Q4" in e for e in result["errors"])

    def test_validate_blank_required(self, data_collector, yunying_template, yunying_answers):
        bad = dict(yunying_answers)
        bad["Q1"] = "   "
        result = data_collector.validate(bad, yunying_template)
        assert result["valid"] is False
        assert result["errors"]

    def test_validate_bad_option_warns(self, data_collector, yunying_template, yunying_answers):
        bad = dict(yunying_answers)
        bad["Q4"] = "每年"  # 不在选项列表
        result = data_collector.validate(bad, yunying_template)
        assert result["valid"] is True  # 选项不合法只是 warning
        assert any("Q4" in w for w in result["warnings"])

    def test_validate_numeric_option_index_ok(self, data_collector, sample_template):
        data = {"Q1": "x", "Q2": "2", "Q3": ""}  # "2" 匹配选项索引 1（每周）
        result = data_collector.validate(data, sample_template)
        assert result["valid"] is True
        assert result["warnings"] == []

    def test_validate_numeric_index_out_of_range_warns(self, data_collector, sample_template):
        data = {"Q1": "x", "Q2": "99"}  # 索引越界 → warning
        result = data_collector.validate(data, sample_template)
        assert result["valid"] is True
        assert any("Q2" in w for w in result["warnings"])

    def test_validate_unknown_qid_warns(self, data_collector, yunying_template, yunying_answers):
        bad = dict(yunying_answers)
        bad["Q99"] = "神秘答案"
        result = data_collector.validate(bad, yunying_template)
        assert any("Q99" in w for w in result["warnings"])


class TestDataCollectorCollectAndExport:
    def test_collect_structured_wrapped(self, data_collector, event_bus, capture_events):
        """存在 answers 键时只取 answers（外层 Q2/junk 被丢弃）；事件 payload 同步。"""
        seen = capture_events(event_bus, "data.collected")
        out = data_collector.collect_structured({"answers": {"Q1": "a"}, "Q2": "b", "junk": "z"})
        assert out == {"Q1": "a"}
        assert seen[0]["data"] == {"Q1": "a"}

    def test_collect_structured_empty_raises(self, data_collector):
        with pytest.raises(CollectionError):
            data_collector.collect_structured({"no_q_keys": 1})

    def test_collect_structured_flat_qkeys(self, data_collector):
        out = data_collector.collect_structured({"Q1": "a", "other": "x"})
        assert out == {"Q1": "a"}

    def test_collect_interactive(self, data_collector, event_bus, capture_events, sample_template, monkeypatch):
        """monkeypatch builtins.input 模拟交互输入。"""
        seen = capture_events(event_bus, "data.collected")
        input_seq = iter(["流程A", "每天", "痛点B"])
        # sample_template: Q1 必填, Q2 必填(选项), Q3 选填
        monkeypatch.setattr("builtins.input", lambda prompt="": next(input_seq))
        answers = data_collector.collect_interactive(sample_template)
        assert answers == {"Q1": "流程A", "Q2": "每天", "Q3": "痛点B"}
        assert seen and seen[0]["answers"] == answers

    def test_collect_interactive_required_empty_raises(
        self, data_collector, sample_template, monkeypatch
    ):
        monkeypatch.setattr("builtins.input", lambda *a, **k: "  ")
        with pytest.raises(CollectionError, match="必填"):
            data_collector.collect_interactive(sample_template)

    def test_export_json(self, data_collector, tmp_path):
        f = tmp_path / "sub" / "out.json"
        ret = data_collector.export({"answers": {"Q1": "a"}}, str(f))
        assert ret == str(f)
        assert json.loads(f.read_text(encoding="utf-8")) == {"answers": {"Q1": "a"}}

    def test_export_markdown(self, data_collector, tmp_path):
        f = tmp_path / "out.md"
        data_collector.export({"answers": {"Q1": "流程", "Q2": "每天"}}, str(f))
        text = f.read_text(encoding="utf-8")
        assert "### Q1" in text and "答案: 流程" in text

    def test_export_docx(self, data_collector, tmp_path):
        from docx import Document

        f = tmp_path / "out.docx"
        data_collector.export({"answers": {"Q1": "流程"}}, str(f))
        assert f.is_file()
        paras = [p.text for p in Document(str(f)).paragraphs]
        assert "答案: 流程" in paras


# ---------------------------------------------------------------------------
# AIAnalysisEngine
# ---------------------------------------------------------------------------


class TestAIAnalysisEngine:
    async def test_analyze_department_returns_parsed_result(
        self, fake_llm_client, event_bus, capture_events, sample_template, sample_answers
    ):
        engine = AIAnalysisEngine(fake_llm_client, event_bus=event_bus)
        seen = capture_events(event_bus, "analysis.done")

        result = await engine.analyze_department(_dept_data("客服部", sample_answers, sample_template))

        assert isinstance(result, AnalysisResult)
        assert result.department == "客服部"
        assert result.confidence == 0.85
        assert result.pain_points == ["手工录入错误率高", "报表汇总依赖人工"]
        assert len(result.opportunities) == 3
        op = result.opportunities[0]
        assert op.process == "工单分派"
        assert op.ai_suggestion == "LLM 自动分类并分派工单"
        assert op.priority == 5
        assert op.difficulty == "中"

        # 事件发布
        assert len(seen) == 1
        assert seen[0]["department"] == "客服部"
        assert seen[0]["result"] is result

        # prompt 包含六维度与答案内容
        msgs, temp, max_tokens = fake_llm_client.calls[0]
        prompt = msgs[0]["content"]
        for dim in ("重复性", "数据密集度", "人力成本", "AI成熟度", "落地难度", "提效预期"):
            assert dim in prompt
        assert sample_answers["Q1"] in prompt
        assert temp == 0.7 and max_tokens == 4096

    async def test_analyze_department_without_template_still_sends_answers(
        self, fake_llm_client, event_bus
    ):
        engine = AIAnalysisEngine(fake_llm_client, event_bus=event_bus)
        result = await engine.analyze_department(_dept_data("设计部", {"Q1": "流程X"}))
        assert result.department == "设计部"
        prompt = fake_llm_client.calls[0][0][0]["content"]
        assert "流程X" in prompt
        assert "### Q1" in prompt

    async def test_analyze_all_concurrent_and_counts(
        self, fake_llm_client, event_bus, sample_template, sample_answers
    ):
        # 三个部门，canned 中每个部门 3 个机会点，其中优先级 5/3/1 → 每个部门 2 个高优(>=3)
        fake_llm_client.delay = 0.05
        engine = AIAnalysisEngine(fake_llm_client, event_bus=event_bus, max_concurrency=3)
        data = [_dept_data(d, sample_answers, sample_template) for d in ("客服部", "财务部", "仓储部")]
        report = await engine.analyze_all(data)

        assert [d.department for d in report.departments] == ["客服部", "财务部", "仓储部"]
        assert report.total_opportunities == 9
        assert report.high_priority_opportunities == 6
        assert fake_llm_client.peak_concurrency >= 2, "3 个部门应并发执行"
        assert len(fake_llm_client.calls) == 3
        assert "3 个部门" in report.overall_summary

    async def test_analyze_all_respects_semaphore_cap(self, fake_llm_client, event_bus):
        """max_concurrency=2 时 4 个部门的峰值并发不超过 2。"""
        fake_llm_client.delay = 0.02
        engine = AIAnalysisEngine(fake_llm_client, event_bus=event_bus, max_concurrency=2)
        data = [_dept_data(d, {"Q1": "x"}) for d in ("d1", "d2", "d3", "d4")]
        await engine.analyze_all(data)
        assert fake_llm_client.peak_concurrency <= 2

    async def test_analyze_all_error_collection(self, fake_llm_client, event_bus):
        """LLM 对某部门抛 LLMError → 该部门记为失败，其他部门成功（gather return_exceptions）。"""

        class FailingClient:
            def __init__(self, fake):
                self._fake = fake

            async def chat(self, messages, temperature=0.7, max_tokens=2048):
                if "失败部" in messages[0]["content"]:
                    raise LLMError("模拟 LLM 失败")
                return await self._fake.chat(messages, temperature, max_tokens)

        fc = FailingClient(fake_llm_client)
        engine = AIAnalysisEngine(fc, event_bus=event_bus)
        data = [_dept_data("成功部", {"Q1": "x"}), _dept_data("失败部", {"Q1": "y"})]
        report = await engine.analyze_all(data)
        assert [d.department for d in report.departments] == ["成功部"]
        assert report.total_opportunities == 3
        assert "分析失败" in report.overall_summary

    async def test_analyze_all_all_failed_summary(self, event_bus):
        from tests.conftest import FakeLLMClient

        failing = FakeLLMClient(fail_with=LLMError("api down"))
        engine = AIAnalysisEngine(failing, event_bus=event_bus)
        report = await engine.analyze_all([_dept_data("客服部", {"Q1": "x"})])
        assert report.departments == []
        assert "所有部门分析失败" in report.overall_summary

    async def test_analyze_department_wraps_llm_error(self, event_bus):
        from tests.conftest import FakeLLMClient

        failing = FakeLLMClient(fail_with=LLMError("api down", code="LLM_ERROR"))
        engine = AIAnalysisEngine(failing, event_bus=event_bus)
        with pytest.raises(AnalysisError, match="分析失败"):
            await engine.analyze_department(_dept_data("客服部", {"Q1": "x"}))

    async def test_parse_llm_response_bad_json_raises(self, event_bus):
        from tests.conftest import FakeLLMClient as _Fake

        bad = _Fake(canned="这不是 JSON {")
        engine = AIAnalysisEngine(bad, event_bus=event_bus)
        with pytest.raises(AnalysisError, match="解析分析结果失败"):
            await engine.analyze_department(_dept_data("客服部", {"Q1": "x"}))

    def test_parse_llm_response_extracts_json_from_markdown(self, event_bus):
        from tests.conftest import FakeLLMClient, SAMPLE_LLM_RESPONSE

        wrapped = "以下是分析结果：\n```json\n" + SAMPLE_LLM_RESPONSE + "\n```"
        fake = FakeLLMClient(canned=wrapped)
        engine = AIAnalysisEngine(fake, event_bus=event_bus)
        result = engine._parse_llm_response(fake.canned, "设计部")
        assert result.department == "设计部"
        assert len(result.opportunities) == 3

    def test_default_concurrency(self, fake_llm_client):
        engine = AIAnalysisEngine(fake_llm_client)
        assert engine._max_concurrency == 3

    def test_build_analysis_prompt_no_answers(self):
        engine = AIAnalysisEngine(None, event_bus=None)
        prompt = engine._build_analysis_prompt("运营部", None, {})
        assert "运营部" in prompt
        assert "重复性分析" in prompt

    def test_generate_overall_summary_no_results(self):
        engine = AIAnalysisEngine(None, event_bus=None)
        assert "所有部门分析失败" in engine._generate_overall_summary([], [])
