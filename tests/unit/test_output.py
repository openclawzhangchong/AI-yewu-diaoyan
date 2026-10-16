"""output 单元测试 — ReportBuilder / SkillPackager。"""

from __future__ import annotations

import json

import openpyxl
import pytest

from infra.event_bus import PACK_DONE, REPORT_GENERATED
from infra.exceptions import PackError, ReportError
from output.report_builder import ReportBuilder
from output.skill_packager import SkillPackager
from tests.conftest import REPO_ROOT


class TestReportBuilderBuild:
    def test_build_json(self, sample_full_report, event_bus, capture_events, tmp_path):
        seen = capture_events(event_bus, REPORT_GENERATED)
        builder = ReportBuilder(event_bus=event_bus, project_root=REPO_ROOT)
        path = builder.build(sample_full_report, "json", tmp_path / "reports")
        assert path.is_file() and path.suffix == ".json"
        data = json.loads(path.read_text(encoding="utf-8"))
        assert data["report_metadata"]["total_departments"] == 2
        assert data["report_metadata"]["total_opportunities"] == 5
        assert data["report_metadata"]["high_priority_opportunities"] == 3
        assert data["departments"][0]["department"] == "客服部"
        assert data["departments"][0]["opportunities"][0]["process"] == "工单分派"

        # 事件 payload 字段
        assert len(seen) == 1
        payload = seen[0]
        assert payload["filepath"] == str(path)
        assert payload["output_path"] == str(path)
        assert payload["format"] == "json"
        assert payload["total_opportunities"] == 5
        assert payload["high_priority_opportunities"] == 3
        assert payload["generated_at"] is not None

    def test_build_markdown_alias(self, sample_full_report, event_bus, capture_events, tmp_path):
        seen = capture_events(event_bus, REPORT_GENERATED)
        builder = ReportBuilder(event_bus=event_bus, project_root=REPO_ROOT)
        path = builder.build(sample_full_report, "markdown", tmp_path)
        assert path.suffix == ".md"
        text = path.read_text(encoding="utf-8")
        assert "AI业务调研提效分析报告" in text
        assert "客服部" in text
        assert "工单分派" in text
        assert "智能数据报表" not in text  # 非样本内容，只应含报告数据
        assert seen[0]["format"] == "markdown"

    def test_build_md_uppercase(self, sample_full_report, event_bus, tmp_path):
        builder = ReportBuilder(event_bus=event_bus, project_root=REPO_ROOT)
        path = builder.build(sample_full_report, "MD", tmp_path)
        assert path.suffix == ".md"

    def test_build_xlsx(self, sample_full_report, event_bus, capture_events, tmp_path):
        seen = capture_events(event_bus, REPORT_GENERATED)
        builder = ReportBuilder(event_bus=event_bus, project_root=REPO_ROOT)
        path = builder.build(sample_full_report, "excel", tmp_path)
        assert path.suffix == ".xlsx"

        wb = openpyxl.load_workbook(str(path))
        assert wb.sheetnames == ["执行摘要", "各部门分析", "机会点明细", "综合建议"]
        ws = wb["机会点明细"]
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0][0] == "部门"
        assert len(rows) == 1 + 5  # 表头 + 5 个机会点
        # 各部门分析 sheet
        ws2 = wb["各部门分析"]
        dept_rows = list(ws2.iter_rows(min_row=2, values_only=True))
        assert [r[0] for r in dept_rows] == ["客服部", "财务部"]
        assert seen[0]["format"] == "excel"

    def test_build_unsupported_format_raises(self, sample_full_report, tmp_path):
        builder = ReportBuilder(project_root=REPO_ROOT)
        with pytest.raises(ReportError) as ei:
            builder.build(sample_full_report, "pdf", tmp_path)
        assert ei.value.code == "REPORT_FORMAT_ERROR"

    def test_build_invalid_report_raises(self, tmp_path):
        builder = ReportBuilder(project_root=REPO_ROOT)
        with pytest.raises(ReportError) as ei:
            builder.build({"not": "a report"}, "json", tmp_path)
        assert ei.value.code == "REPORT_INVALID_INPUT"

    def test_build_creates_nested_dir(self, sample_full_report, tmp_path):
        builder = ReportBuilder(project_root=REPO_ROOT)
        deep = tmp_path / "a" / "b" / "c"
        path = builder.build(sample_full_report, "json", deep)
        assert path.parent == deep

    def test_generator_property(self, sample_full_report):
        builder = ReportBuilder(project_root=REPO_ROOT)
        assert builder.generator is not None
        assert builder.event_bus is not None


class TestSkillPackager:
    @pytest.fixture
    def packager(self, event_bus):
        return SkillPackager(project_root=REPO_ROOT, event_bus=event_bus)

    def test_pack_light_structure(self, packager, event_bus, capture_events, tmp_path):
        seen = capture_events(event_bus, PACK_DONE)
        out = packager.pack_light(tmp_path / "skill_light")
        assert out.is_dir()
        # SKILL.md 存在且非空
        skill_md = out / "SKILL.md"
        assert skill_md.is_file()
        assert "ai-efficiency-analyzer" in skill_md.read_text(encoding="utf-8")
        # manifest.json 字段
        manifest = json.loads((out / "manifest.json").read_text(encoding="utf-8"))
        assert manifest["name"] == "ai-efficiency-analyzer"
        assert manifest["version"] == "0.1.0"
        assert manifest["author"]
        assert manifest["category"] == "business-analysis"
        assert manifest["entry_point"] == "entry.main:main"
        assert isinstance(manifest["dependencies"], list)
        assert "generated_at" in manifest
        # prompts/ 目录存在（repo 无 prompts/ 时保留空目录结构）
        assert (out / "prompts").is_dir()
        # 5 个部门模板 JSON 被复制
        tpl_jsons = sorted(p.name for p in (out / "templates").glob("*.json"))
        assert len(tpl_jsons) == 5
        for name in tpl_jsons:
            json.loads((out / "templates" / name).read_text(encoding="utf-8"))
        # 轻量模式不含代码目录
        assert not (out / "code").exists()
        # pack.done 事件
        assert seen[0]["mode"] == "light"
        assert seen[0]["skill_dir"] == str(out)
        assert seen[0]["name"] == "ai-efficiency-analyzer"

    def test_pack_full_structure(self, packager, event_bus, capture_events, tmp_path):
        seen = capture_events(event_bus, PACK_DONE)
        out = packager.pack_full(tmp_path / "skill_full")
        for code_dir in ("domain", "core", "infra", "entry", "output"):
            assert (out / "code" / code_dir).is_dir(), f"full 模式应包含 code/{code_dir}"
        assert (out / "code" / "core" / "ai_engine.py").is_file()
        # 依赖清单
        req = out / "requirements.txt"
        assert req.is_file()
        assert "openpyxl" in req.read_text(encoding="utf-8")
        # 轻量部分也在
        assert (out / "SKILL.md").is_file()
        assert (out / "manifest.json").is_file()
        # 打包产物不含 __pycache__
        assert not list((out / "code").rglob("__pycache__"))
        assert seen[0]["mode"] == "full"

    def test_validate_pass_after_pack(self, packager, tmp_path):
        out = packager.pack_light(tmp_path / "v")
        report = packager.validate(out)
        assert report["valid"] is True
        assert report["errors"] == []
        assert report["files_checked"] >= 7  # SKILL.md + manifest + 5 模板

    def test_validate_broken_manifest_raises_pack_error(self, packager, tmp_path):
        out = packager.pack_light(tmp_path / "broken")
        # 破坏 manifest：删掉必需字段
        manifest = json.loads((out / "manifest.json").read_text(encoding="utf-8"))
        del manifest["entry_point"]
        (out / "manifest.json").write_text(
            json.dumps(manifest, ensure_ascii=False), encoding="utf-8"
        )
        with pytest.raises(PackError) as ei:
            packager.validate(out)
        assert ei.value.code == "PACK_VALIDATE_FAILED"
        assert "entry_point" in str(ei.value)

    def test_validate_broken_template_json_raises(self, packager, tmp_path):
        out = packager.pack_light(tmp_path / "broken_tpl")
        tpl_files = list((out / "templates").glob("*.json"))
        assert tpl_files
        tpl_files[0].write_text("{not valid json", encoding="utf-8")
        with pytest.raises(PackError, match="模板 JSON 解析失败"):
            packager.validate(out)

    def test_validate_missing_skill_md_raises(self, packager, tmp_path):
        out = packager.pack_light(tmp_path / "no_md")
        (out / "SKILL.md").unlink()
        with pytest.raises(PackError, match="SKILL.md 缺失"):
            packager.validate(out)

    def test_validate_nonexistent_dir_raises(self, packager, tmp_path):
        with pytest.raises(PackError, match="目录不存在"):
            packager.validate(tmp_path / "not_there")

    def test_pack_light_publishes_even_with_custom_bus(self, event_bus, tmp_path):
        seen = []
        event_bus.subscribe(PACK_DONE, seen.append)
        p = SkillPackager(project_root=REPO_ROOT, event_bus=event_bus)
        p.pack_light(tmp_path / "ev")
        assert len(seen) == 1

    def test_generate_manifest_alone(self, packager, tmp_path):
        out = tmp_path / "only_manifest"
        p = packager.generate_manifest(out)
        assert p.is_file()
        manifest = json.loads(p.read_text(encoding="utf-8"))
        assert set(manifest) >= {"name", "version", "author", "category", "entry_point", "dependencies"}
