"""core.report_generator 单元测试 — ReportGenerator（真实 jinja 模板 + 纯函数）。"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

from core.report_generator import ReportGenerator
from domain.analysis import AnalysisResult, FullAnalysisReport


@pytest.fixture
def generator(report_generator) -> ReportGenerator:
    return report_generator


class TestGenerateMarkdown:
    def test_contains_report_sections(self, generator, sample_full_report):
        md = generator.generate_markdown(sample_full_report)
        assert "AI业务调研提效分析报告" in md
        assert "执行摘要" in md
        assert "客服部" in md and "财务部" in md
        assert "工单分派" in md  # 机会点 process
        assert "综合建议" in md
        # 机会点字段渲染
        assert "LLM 自动分类并分派工单" in md

    def test_empty_report_renders(self, generator):
        md = generator.generate_markdown(FullAnalysisReport())
        assert "执行摘要" in md
        assert "分析部门数: 0" in md

    def test_format_percent_filter_applied(self, generator, sample_analysis_result):
        """jinja 自定义过滤器 format_percent 渲染置信度为百分比（0.85 → 85%）。"""
        md = generator.generate_markdown(
            FullAnalysisReport(departments=[sample_analysis_result])
        )
        assert "85%" in md


class TestGenerateExcel:
    def test_excel_structure(self, generator, sample_full_report, tmp_path):
        out = tmp_path / "r.xlsx"
        ret = generator.generate_excel(sample_full_report, str(out))
        assert ret == str(out)
        assert out.is_file()

        wb = openpyxl.load_workbook(str(out))
        assert wb.sheetnames == ["执行摘要", "各部门分析", "机会点明细", "综合建议"]

        ws = wb["机会点明细"]
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0][0] == "部门"
        assert len(rows) == 1 + 5  # 表头 + 5 个机会点

        ws2 = wb["各部门分析"]
        dept_rows = list(ws2.iter_rows(min_row=2, values_only=True))
        assert [r[0] for r in dept_rows] == ["客服部", "财务部"]

    def test_creates_parent_dirs(self, generator, sample_full_report, tmp_path):
        deep = tmp_path / "a" / "b" / "c.xlsx"
        generator.generate_excel(sample_full_report, str(deep))
        assert deep.is_file()


class TestGenerateJson:
    def test_json_structure(self, generator, sample_full_report, tmp_path):
        out = tmp_path / "r.json"
        generator.generate_json(sample_full_report, str(out))
        data = json.loads(out.read_text(encoding="utf-8"))

        assert data["report_metadata"]["total_departments"] == 2
        assert data["report_metadata"]["total_opportunities"] == 5
        assert data["report_metadata"]["high_priority_opportunities"] == 3
        assert data["departments"][0]["department"] == "客服部"
        assert data["departments"][0]["opportunities"][0]["process"] == "工单分派"
        assert data["departments"][0]["confidence"] == 0.85

    def test_generated_at_defaults_when_missing(self, generator, tmp_path):
        rep = FullAnalysisReport()  # generated_at=None → 现生成时间
        out = tmp_path / "r2.json"
        generator.generate_json(rep, str(out))
        data = json.loads(out.read_text(encoding="utf-8"))
        assert data["report_metadata"]["generated_at"]


class TestGenerateSummary:
    def test_summary_text(self, generator, sample_full_report):
        text = generator.generate_summary(sample_full_report)
        assert "AI业务调研提效分析报告" in text
        assert "客服部" in text and "财务部" in text
        assert "总机会点数: 5" in text
        assert "高优先级机会点: 3" in text

    def test_summary_confidence_percent(self, generator, sample_analysis_result):
        text = generator.generate_summary(
            FullAnalysisReport(departments=[sample_analysis_result])
        )
        assert "85%" in text


class TestGenerateDepartmentReport:
    def test_sections(self, generator, sample_analysis_result):
        dr = generator.generate_department_report(sample_analysis_result)
        assert dr.department == "客服部"
        assert dr.executive_summary == sample_analysis_result.summary
        titles = [s.title for s in dr.sections]
        assert titles == ["执行摘要", "痛点分析", "AI提效机会", "建议方案"]
        # 建议方案：优先级 5 → 短期优先推进
        assert "短期优先推进" in dr.recommendation

    def test_recommendation_no_opportunities(self, generator):
        r = AnalysisResult(department="空部", process="x")
        dr = generator.generate_department_report(r)
        assert "暂无明显的AI提效机会" in dr.recommendation


class TestPureHelpers:
    def test_format_number(self, generator):
        assert generator._format_number(1234.6) == "1,235"
        assert generator._format_number("abc") == "abc"

    def test_format_percent(self, generator):
        assert generator._format_percent(0.85) == "85%"
        assert generator._format_percent("x") == "x"

    def test_format_datetime(self, generator):
        dt = datetime(2026, 10, 16, 8, 0, 0)
        assert generator._format_datetime(dt) == "2026-10-16 08:00:00"
        assert generator._format_datetime("now") == "now"

    def test_truncate(self, generator):
        assert generator._truncate("short") == "short"
        assert generator._truncate("x" * 200, 100).endswith("...")
        assert len(generator._truncate("x" * 200, 100)) == 103
