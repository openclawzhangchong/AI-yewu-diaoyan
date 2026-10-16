"""集成 E2E — fixture 文件 → DataCollector.parse_file → 假引擎分析 → ReportBuilder 生成。

全程离线：LLM 用 FakeLLMClient 桩；ReportBuilder 用真实 ReportGenerator
（jinja 模板指向 repo 的 templates/）。
"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from core.ai_engine import AIAnalysisEngine
from core.data_collector import DataCollector
from infra.event_bus import ANALYSIS_DONE, EventBus, REPORT_GENERATED
from output.report_builder import ReportBuilder
from tests.conftest import FakeLLMClient, REPO_ROOT


@pytest.fixture
def e2e_pipeline(md_fixture_path, event_bus, fake_llm_client, capture_events):
    """组装 parse → analyze → build 的离线管线，返回各环节产物。"""
    collector = DataCollector(event_bus=event_bus)
    engine = AIAnalysisEngine(fake_llm_client, event_bus=event_bus)
    builder = ReportBuilder(event_bus=event_bus, project_root=REPO_ROOT)
    analysis_events = capture_events(event_bus, ANALYSIS_DONE)
    report_events = capture_events(event_bus, REPORT_GENERATED)
    return {
        "collector": collector,
        "engine": engine,
        "builder": builder,
        "analysis_events": analysis_events,
        "report_events": report_events,
        "input": md_fixture_path,
    }


class TestE2E:
    async def test_fixture_file_to_report_json(
        self, e2e_pipeline, tmp_path, json_fixture_path, txt_fixture_path
    ):
        input_files = [e2e_pipeline["input"], txt_fixture_path, json_fixture_path]

        # 1. 解析各格式文件
        collector = e2e_pipeline["collector"]
        dept_data = []
        for f in input_files:
            parsed = collector.parse_file(str(f))
            assert parsed.get("answers"), f"{f.name} 应解析出答案"
            dept_data.append({"department": f.stem, "answers": parsed["answers"]})

        # 2. 并发分析（假 LLM，无网络）
        report = await e2e_pipeline["engine"].analyze_all(dept_data)
        assert [d.department for d in report.departments] == [
            "survey_kefu", "survey_caiwu", "survey_cangchu"
        ]
        assert report.total_opportunities == 9  # 每部门 3 个
        assert report.high_priority_opportunities == 6  # 每部门 2 个 (优先级 5、3)

        # analysis.done 事件每个部门一次
        assert len(e2e_pipeline["analysis_events"]) == 3
        assert all("result" in ev for ev in e2e_pipeline["analysis_events"])

        # 3. ReportBuilder 生成 json 报告
        out_dir = tmp_path / "e2e_reports"
        path = e2e_pipeline["builder"].build(report, "json", out_dir)
        assert path.is_file() and path.suffix == ".json"
        data = json.loads(path.read_text(encoding="utf-8"))
        assert data["report_metadata"]["total_departments"] == 3
        assert data["report_metadata"]["total_opportunities"] == 9
        assert data["report_metadata"]["high_priority_opportunities"] == 6
        dept_names = {d["department"] for d in data["departments"]}
        assert dept_names == {"survey_kefu", "survey_caiwu", "survey_cangchu"}
        # 机会点字段完整
        first_opp = data["departments"][0]["opportunities"][0]
        assert set(first_opp) == {
            "process", "current_process", "ai_suggestion",
            "expected_efficiency_gain", "difficulty", "priority",
            "expected_roi", "timeline",
        }

        # report.generated 事件 payload
        assert len(e2e_pipeline["report_events"]) == 1
        payload = e2e_pipeline["report_events"][0]
        assert payload["filepath"] == str(path)
        assert payload["format"] == "json"
        assert payload["total_opportunities"] == 9
        assert payload["high_priority_opportunities"] == 6
        assert payload["generated_at"] is not None

    async def test_fixture_file_to_markdown_report(self, e2e_pipeline, tmp_path):
        parsed = e2e_pipeline["collector"].parse_file(str(e2e_pipeline["input"]))
        result = await e2e_pipeline["engine"].analyze_department(
            {"department": "客服部", "answers": parsed["answers"]}
        )
        from domain.analysis import FullAnalysisReport

        full = FullAnalysisReport(
            departments=[result],
            total_opportunities=len(result.opportunities),
            high_priority_opportunities=sum(1 for o in result.opportunities if o.priority >= 3),
            overall_summary=result.summary,
            generated_at="2026-10-16 12:00:00",
        )
        path = e2e_pipeline["builder"].build(full, "markdown", tmp_path / "md")
        text = path.read_text(encoding="utf-8")
        assert "AI业务调研提效分析报告" in text
        assert "工单分派" in text
        assert "客服部" in text

    async def test_fixture_file_to_xlsx_report(self, e2e_pipeline, tmp_path):
        parsed = e2e_pipeline["collector"].parse_file(str(e2e_pipeline["input"]))
        result = await e2e_pipeline["engine"].analyze_department(
            {"department": "客服部", "answers": parsed["answers"]}
        )
        from domain.analysis import FullAnalysisReport

        full = FullAnalysisReport(
            departments=[result],
            total_opportunities=len(result.opportunities),
            high_priority_opportunities=sum(1 for o in result.opportunities if o.priority >= 3),
            overall_summary=result.summary,
        )
        path = e2e_pipeline["builder"].build(full, "excel", tmp_path / "xlsx")
        wb = openpyxl.load_workbook(str(path))
        assert wb.sheetnames[0] == "执行摘要"
        ws = wb["机会点明细"]
        assert len(list(ws.iter_rows(values_only=True))) == 1 + 3


class TestEventBusAcrossLayers:
    """各层组件共享同一 EventBus 时的端到端事件流。"""

    async def test_all_events_published(self, event_bus, capture_events, fake_llm_client, tmp_path):
        from core.data_collector import DataCollector

        seen_analysis = capture_events(event_bus, ANALYSIS_DONE)
        seen_report = capture_events(event_bus, REPORT_GENERATED)

        collector = DataCollector(event_bus=event_bus)
        engine = AIAnalysisEngine(fake_llm_client, event_bus=event_bus)
        builder = ReportBuilder(event_bus=event_bus, project_root=REPO_ROOT)

        parsed = collector.parse_file(str(REPO_ROOT / "tests" / "fixtures" / "survey_kefu.md"))
        result = await engine.analyze_department({"department": "客服部", "answers": parsed["answers"]})
        from domain.analysis import FullAnalysisReport

        full = FullAnalysisReport(departments=[result])
        builder.build(full, "json", tmp_path)

        assert len(seen_analysis) == 1
        assert len(seen_report) == 1
